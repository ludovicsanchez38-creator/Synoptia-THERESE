"""B-797 (cycle 9) : la conversion `int(cell)` du repli Markdown mangeait le zéro
initial d'un code postal (« 04100 » devenait 4100)."""

from __future__ import annotations

from pathlib import Path

import pytest
from app.services.skills.xlsx_generator import XlsxSkill

TABLEAU = """| Ville | Code postal | Effectif |
|---|---|---|
| Manosque | 04100 | 12 |
| Digne | 04000 | 7 |
"""


def test_un_nombre_a_zero_initial_reste_un_texte(tmp_path) -> None:
    skill = XlsxSkill(tmp_path)
    tables = skill._parse_markdown_tables(TABLEAU)
    assert tables, "aucun tableau reconnu"
    lignes = tables[0]["rows"]
    assert lignes[0][1] == "04100"
    assert lignes[1][1] == "04000"
    assert lignes[0][2] == 12


@pytest.mark.asyncio
async def test_le_classeur_livre_garde_le_code_postal_en_texte(tmp_path) -> None:
    """Relecture T2 (c9) : le test ci-dessus s'arrêtait à l'analyseur. Ici on lit
    la cellule réellement écrite dans le fichier produit."""
    from app.services.skills.base import SkillParams
    from openpyxl import load_workbook

    skill = XlsxSkill(tmp_path)
    resultat = await skill.execute(SkillParams(title="Effectifs", content=TABLEAU))
    fichier = Path(resultat.file_path)
    wb = load_workbook(fichier, read_only=True)
    try:
        ws = wb.active
        assert ws is not None
        valeurs = [c.value for row in ws.iter_rows(max_row=10) for c in row]
    finally:
        wb.close()
    assert "04100" in valeurs, valeurs
    assert 4100 not in valeurs
