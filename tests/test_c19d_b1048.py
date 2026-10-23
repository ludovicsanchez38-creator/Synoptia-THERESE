from app.services.skills.pptx_generator import PptxSkill
from pptx import Presentation


def test_signature_du_repli_accentuee(tmp_path):
    presentation = Presentation()
    PptxSkill(output_dir=tmp_path)._add_end_slide(presentation)
    assert any('Généré par THÉRÈSE - Synoptïa' in shape.text
               for shape in presentation.slides[-1].shapes if shape.has_text_frame)
