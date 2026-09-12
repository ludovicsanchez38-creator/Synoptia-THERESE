"""Non-régression du socle, vérifiée par le COMPORTEMENT (B-039, B-332).

Remplace les gardes textuelles de `test_regression.py` sur : le tueur de
zombies (BUG-007), CORS et l'accès réseau privé (BUG-022), l'initialisation
paresseuse du chiffrement (BUG-013) et la clé de secours (BUG-050), les
schémas (BUG-024, 044, 051, F-17, skill_id), le profil (BUG-026),
l'extraction XLSX, l'envoi de fichiers, les coûts (BUG-028), les clés
corrompues (BUG-051), le préfixe des clés Gemini (BUG-099), l'indexation des
contacts CRM (BUG-102), les identifiants CRM (F-13), le chemin des outils MCP
(BUG-062), les délais MCP (BUG-062b), les consignes des générateurs Office,
nb_slides, la version, les migrations de colonnes (P0-IA-3, BUG-130), l'outil
agenda (P0-PROD-3, QW2) et l'ordre des messages (BUG-031).
"""

from __future__ import annotations

import inspect
import os
import re
import sqlite3
import subprocess
import sys
import types
from contextlib import closing
from datetime import UTC, datetime
from pathlib import Path
from unittest.mock import AsyncMock

import pytest

# ---------------------------------------------------------------- point d'entrée


class TestTueurDeZombies:
    def test_il_epargne_le_processus_courant_et_son_parent(self, monkeypatch):
        import os

        import main as point_d_entree  # src/backend/main.py, sans effet de bord hors gel PyInstaller

        moi, parent = os.getpid(), os.getppid()
        tues: list[int] = []

        def faux_run(commande, **kwargs):
            if commande[0] == "pgrep":
                return types.SimpleNamespace(stdout=f"{moi}\n{parent}\n424242\n", returncode=0)
            return types.SimpleNamespace(stdout="", returncode=0)

        def faux_kill(pid, sig):
            tues.append(pid)

        monkeypatch.setattr(subprocess, "run", faux_run)
        monkeypatch.setattr(os, "kill", faux_kill)
        monkeypatch.setattr(sys, "platform", "darwin")

        point_d_entree._kill_zombie_backends()

        assert 424242 in tues, "l'ancien backend doit être arrêté"
        assert moi not in tues and parent not in tues, "ni soi-même ni le bootloader PyInstaller (BUG-007)"


# ---------------------------------------------------------------- CORS et réseau privé


def _sync(client):
    """Le client de test asynchrone n'expose pas OPTIONS : on passe par le TestClient qu'il enveloppe."""
    for nom in ("tc", "_tc", "client", "_client"):
        if hasattr(client, nom):
            return getattr(client, nom)
    pytest.fail("TestClient sous-jacent introuvable")


class TestCorsEtReseauPrive:
    @pytest.mark.asyncio
    @pytest.mark.parametrize("origine", ["tauri://localhost", "https://tauri.localhost", "http://tauri.localhost"])
    async def test_les_trois_origines_tauri_sont_admises(self, client, origine):
        reponse = _sync(client).options(
            "/api/health",
            headers={"Origin": origine, "Access-Control-Request-Method": "GET"},
        )
        assert reponse.headers.get("access-control-allow-origin") == origine, (origine, dict(reponse.headers))

    @pytest.mark.asyncio
    async def test_le_pre_vol_reseau_prive_est_accorde(self, client):
        reponse = _sync(client).options(
            "/api/health",
            headers={
                "Origin": "http://tauri.localhost",
                "Access-Control-Request-Method": "GET",
                "Access-Control-Request-Private-Network": "true",
            },
        )
        assert reponse.headers.get("access-control-allow-private-network") == "true", (
            "WebView2 143+ bloque http://tauri.localhost -> 127.0.0.1 sans cet en-tête (BUG-022)"
        )


# ---------------------------------------------------------------- chiffrement


class FauxTrousseau:
    """Module keyring factice : un secret en mémoire, et le journal des écritures."""

    def __init__(self, secret: str | None = None):
        self.secret = secret
        self.ecrits: list[str] = []

    def get_password(self, service, compte):
        return self.secret

    def set_password(self, service, compte, valeur):
        self.secret = valeur
        self.ecrits.append(valeur)

    def delete_password(self, service, compte):
        self.secret = None


@pytest.fixture
def chiffrement(monkeypatch, tmp_path):
    from app.services import encryption as module

    module.EncryptionService._instance = None
    module.EncryptionService._fernet = None
    module.EncryptionService._using_keychain = False
    monkeypatch.setattr(module, "KEY_FILE", tmp_path / ".encryption_key")
    monkeypatch.setattr(module, "_try_keyring_available", lambda: True)
    trousseau = FauxTrousseau()
    monkeypatch.setitem(sys.modules, "keyring", trousseau)
    yield module, trousseau
    module.EncryptionService._instance = None
    module.EncryptionService._fernet = None
    module.EncryptionService._using_keychain = False


class TestChiffrementParesseux:
    def test_construire_le_service_ne_touche_pas_au_trousseau_mais_chiffrer_oui(self, chiffrement):
        module, trousseau = chiffrement
        appels: list[str] = []
        trousseau.get_password = lambda s, c: appels.append("lecture") or None  # type: ignore[method-assign]

        service = module.EncryptionService()
        assert service._fernet is None and appels == [], "le Keychain ne doit pas être ouvert au démarrage (BUG-013)"

        chiffre = service.encrypt("secret")
        assert appels, "le premier chiffrement déclenche l'initialisation"
        assert service.decrypt(chiffre) == "secret"


class TestCleDeSecours:
    def test_une_cle_trousseau_est_recopiee_dans_le_fichier_de_secours(self, chiffrement):
        module, trousseau = chiffrement
        trousseau.secret = "cle-trousseau-ABC"
        cle = module.EncryptionService()._get_or_create_key()
        assert cle == b"cle-trousseau-ABC"
        assert module.KEY_FILE.read_bytes() == b"cle-trousseau-ABC", "sans fichier de secours, une signature binaire changée perd tout (BUG-050)"

    def test_quand_les_deux_divergent_le_fichier_gagne_et_repare_le_trousseau(self, chiffrement):
        module, trousseau = chiffrement
        module.KEY_FILE.write_bytes(b"cle-fichier-ANCIENNE")
        trousseau.secret = "cle-trousseau-REGENEREE"
        cle = module.EncryptionService()._get_or_create_key()
        assert cle == b"cle-fichier-ANCIENNE", "la clé du fichier est celle qui a chiffré les données existantes"
        assert trousseau.ecrits == ["cle-fichier-ANCIENNE"], "le trousseau doit être remis en accord"

    def test_la_migration_vers_le_trousseau_garde_le_fichier(self, chiffrement):
        module, trousseau = chiffrement
        module.KEY_FILE.write_bytes(b"cle-fichier-SEULE")
        cle = module.EncryptionService()._get_or_create_key()
        assert cle == b"cle-fichier-SEULE" and trousseau.secret == "cle-fichier-SEULE"
        assert module.KEY_FILE.exists(), "le fichier reste une sauvegarde après migration"


class TestB748ProfilTemporaireSansTrousseau:
    def test_un_data_dir_temporaire_necrit_jamais_dans_le_trousseau(
        self, monkeypatch, tmp_path
    ):
        from app.services import encryption as module

        profil_temporaire = tmp_path / "profil-isole"
        fichier_cle = profil_temporaire / ".encryption_key"
        trousseau = FauxTrousseau()

        assert module._compte_trousseau(Path.home() / ".therese") == "encryption-key", (
            "le profil réel doit conserver le compte Keychain historique"
        )

        monkeypatch.setenv("THERESE_DATA_DIR", str(profil_temporaire))
        monkeypatch.setattr(module._settings, "data_dir", profil_temporaire)
        monkeypatch.setattr(module, "THERESE_DIR", profil_temporaire)
        monkeypatch.setattr(module, "KEY_FILE", fichier_cle)
        monkeypatch.setattr(
            module, "KEYCHAIN_ACCOUNT", module._compte_trousseau(profil_temporaire)
        )
        monkeypatch.setattr(module, "_try_keyring_available", lambda: True)
        monkeypatch.setitem(sys.modules, "keyring", trousseau)
        module.EncryptionService._instance = None
        module.EncryptionService._fernet = None
        module.EncryptionService._using_keychain = False

        try:
            service = module.EncryptionService()
            chiffre = service.encrypt("secret temporaire")

            assert service.decrypt(chiffre) == "secret temporaire"
            assert fichier_cle.exists(), (
                "un profil temporaire doit créer sa clé locale .encryption_key"
            )
            assert trousseau.ecrits == [], (
                "THERESE_DATA_DIR temporaire ne doit jamais appeler keyring.set_password"
            )
        finally:
            module.EncryptionService._instance = None
            module.EncryptionService._fernet = None
            module.EncryptionService._using_keychain = False


# ---------------------------------------------------------------- schémas


class TestSchemas:
    def test_le_profil_porte_les_champs_de_facturation(self):
        from app.models.schemas import UserProfileResponse, UserProfileUpdate

        for schema in (UserProfileUpdate, UserProfileResponse):
            assert {"address", "siren", "tva_intra"} <= set(schema.model_fields), schema.__name__

    def test_la_requete_de_chat_accepte_un_skill_et_des_fichiers(self):
        from app.models.schemas import ChatRequest

        assert {"skill_id", "file_paths"} <= set(ChatRequest.model_fields)

    def test_la_configuration_expose_les_cles_corrompues(self):
        from app.models.schemas import ConfigResponse

        assert "corrupted_keys" in ConfigResponse.model_fields

    def test_le_board_connait_le_mode_souverain(self):
        from app.models.board import BoardMode, BoardRequest

        assert {m.value for m in BoardMode} >= {"cloud", "sovereign"}
        assert BoardRequest.model_fields["mode"].default == BoardMode.CLOUD
        assert "ollama_models" in BoardRequest.model_fields

    def test_le_fournisseur_d_images_fal_existe(self):
        from app.services.image_generator import ImageGeneratorService, ImageProvider

        assert ImageProvider.FAL.value == "fal-flux-pro"
        assert callable(getattr(ImageGeneratorService, "_generate_fal", None))


# ---------------------------------------------------------------- profil, fichiers, coûts


class TestProfilEtFichiers:
    @pytest.mark.asyncio
    async def test_le_profil_est_range_dans_sa_categorie_et_s_efface_avec_elle(self, db_session, monkeypatch):
        from app.models.entities import Preference
        from app.services import user_profile as module
        from sqlalchemy import select

        async def sans_indexation(profile):
            return None

        monkeypatch.setattr(module, "_embed_profile", sans_indexation)
        await module.set_user_profile(db_session, module.UserProfile(name="Jérôme"))
        lignes = (await db_session.execute(select(Preference).where(Preference.category == module.PROFILE_CATEGORY))).scalars().all()
        assert lignes, "le profil doit être rangé sous PROFILE_CATEGORY (BUG-026)"

        assert await module.delete_user_profile(db_session) is True
        restantes = (await db_session.execute(select(Preference).where(Preference.category == module.PROFILE_CATEGORY))).scalars().all()
        assert restantes == []

    def test_un_classeur_xlsx_est_lu(self, tmp_path):
        import openpyxl
        from app.services.file_parser import extract_text

        chemin = tmp_path / "budget.xlsx"
        classeur = openpyxl.Workbook()
        classeur.active["A1"] = "Poste"
        classeur.active["B1"] = "Chauffage cellule 7"
        classeur.save(chemin)
        texte = extract_text(chemin) or ""
        assert "Chauffage cellule 7" in texte

    @pytest.mark.asyncio
    async def test_l_envoi_d_un_fichier_refuse_les_extensions_dangereuses(self, client, monkeypatch):
        from app.services import indexation

        monkeypatch.setattr(indexation, "get_qdrant_service", lambda: AsyncMock())
        projet = await client.post("/api/memory/projects", json={"name": "Dossier upload", "description": "x"})
        assert projet.status_code == 200, projet.text
        champs = {"project_id": projet.json()["id"]}
        refus = await client.post("/api/files/upload", data=champs, files={"file": ("virus.exe", b"MZ", "application/octet-stream")})
        assert refus.status_code in (400, 415, 422), refus.text
        assert "exe" in refus.text or "extension" in refus.text.lower()
        ok = await client.post("/api/files/upload", data=champs, files={"file": ("note.txt", b"Bonjour", "text/plain")})
        assert ok.status_code == 200, ok.text


class TestCouts:
    def test_le_prefixe_du_fournisseur_n_empeche_pas_de_trouver_le_tarif(self):
        from app.services.token_tracker import TokenTracker

        suivi = TokenTracker.__new__(TokenTracker)
        nu = suivi.estimate_cost("claude-sonnet-4-6", 1_000_000, 1_000_000)
        prefixe = suivi.estimate_cost("anthropic/claude-sonnet-4-6", 1_000_000, 1_000_000)
        inconnu = suivi.estimate_cost("modele-inexistant-zz", 1_000_000, 1_000_000)
        assert prefixe == nu, "le préfixe fournisseur (OpenRouter) doit être ignoré (BUG-028)"
        assert nu != inconnu, "un modèle connu ne doit pas retomber sur le tarif par défaut"


# ---------------------------------------------------------------- clés API


class TestClesApi:
    @pytest.mark.asyncio
    async def test_une_cle_illisible_est_signalee_comme_corrompue(self, client, db_session):
        from app.models.entities import Preference
        from app.services.encryption import encrypt_value

        jeton = encrypt_value("sk-valide")
        abime = jeton[:-6] + "AAAAA="  # base64 correct, signature fausse : indéchiffrable
        db_session.add(Preference(key="openai_api_key", value=abime, category="llm"))
        await db_session.commit()
        reponse = await client.get("/api/config/")
        assert reponse.status_code == 200, reponse.text
        assert "openai" in reponse.json().get("corrupted_keys", []), reponse.json().get("corrupted_keys")

    @pytest.mark.asyncio
    async def test_une_cle_gemini_recente_sans_prefixe_aiza_est_acceptee(self, client):
        reponse = await client.post("/api/config/api-key", json={"provider": "gemini", "api_key": "AQ.abcdefghijklmnopqrstuvwxyz0123456789"})
        assert reponse.status_code == 200, reponse.text


# ---------------------------------------------------------------- CRM


class TestCrm:
    @pytest.mark.asyncio
    async def test_creer_et_modifier_un_contact_crm_reindexe_la_memoire(self, client, monkeypatch):
        from app.routers import crm as crm_module
        from app.routers import memory as memory_module

        embed = AsyncMock()
        monkeypatch.setattr(memory_module, "_embed_contact", embed)
        monkeypatch.setattr(crm_module, "_embed_contact", embed, raising=False)

        creation = await client.post(
            "/api/crm/contacts", json={"first_name": "Ana", "last_name": "Ruiz", "email": "ana@exemple.fr", "company": "Ruiz SARL"}
        )
        assert creation.status_code == 200, creation.text
        assert embed.await_count >= 1, "la création CRM doit indexer le contact (BUG-102)"

        modification = await client.patch(f"/api/memory/contacts/{creation.json()['id']}", json={"company": "Ruiz & Fils"})
        assert modification.status_code == 200, modification.text
        assert embed.await_count >= 2, "la modification doit réindexer"

    @pytest.mark.asyncio
    async def test_les_identifiants_google_sont_verifies(self, client):
        reponse = await client.post("/api/crm/sync/credentials", json={"client_id": "pas-un-client", "client_secret": "x"})
        assert reponse.status_code == 400, reponse.text
        assert "apps.googleusercontent.com" in reponse.text


# ---------------------------------------------------------------- MCP


class TestMcp:
    def test_le_chemin_enrichi_trouve_node_hors_du_path_de_l_application(self, monkeypatch, tmp_path):
        from app.services.mcp_service import build_mcp_enriched_path

        (tmp_path / ".nvm" / "versions" / "node" / "v22.19.0" / "bin").mkdir(parents=True)
        (tmp_path / ".volta" / "bin").mkdir(parents=True)
        monkeypatch.setenv("HOME", str(tmp_path))
        # Windows : expanduser lit USERPROFILE avant HOME.
        monkeypatch.setenv("USERPROFILE", str(tmp_path))
        monkeypatch.setattr(Path, "home", classmethod(lambda cls: tmp_path))
        chemin = build_mcp_enriched_path()
        # Seuls les dossiers qui existent entrent dans le PATH : Homebrew n'est
        # attendu que là où il est installé (rouge sur ubuntu et Windows le
        # 08/09/2026, où ce test réécrit supposait macOS). Les segments sont
        # comparés normalisés : le service assemble « ~/.volta/bin » avec une
        # barre oblique, Windows le relit avec des antislashs.
        segments = {os.path.normpath(p) for p in chemin.split(os.pathsep)}
        if Path("/opt/homebrew/bin").is_dir():
            assert os.path.normpath("/opt/homebrew/bin") in segments
        assert os.path.normpath(str(tmp_path / ".volta" / "bin")) in segments, chemin
        assert os.path.normpath(str(tmp_path / ".nvm" / "versions" / "node" / "v22.19.0" / "bin")) in segments, chemin

    def test_l_appel_d_un_outil_dispose_d_au_moins_une_minute(self):
        from app.services.mcp_service import MCPService

        delais = [
            signature.parameters["timeout"].default
            for _nom, fonction in inspect.getmembers(MCPService, inspect.isfunction)
            for signature in [inspect.signature(fonction)]
            if "timeout" in signature.parameters and signature.parameters["timeout"].default is not inspect._empty
        ]
        assert delais and min(delais) >= 60.0, f"un outil MCP lent dépasse 30 s (BUG-062b) : {delais}"


# ---------------------------------------------------------------- générateurs Office


class TestConsignesDesGenerateurs:
    def test_docx_interdit_le_recapitulatif_et_signe_avec_les_accents(self, tmp_path):
        from app.services.skills.docx_generator import DocxSkill

        consigne = DocxSkill(output_dir=tmp_path).get_system_prompt_addition()
        assert "récap" in consigne.lower()
        assert "Synoptïa" in consigne

    def test_pptx_interdit_le_markdown_et_impose_le_nombre_de_diapositives(self, tmp_path):
        from app.services.skills.pptx_generator import PptxSkill

        consigne = PptxSkill(output_dir=tmp_path).get_system_prompt_addition()
        assert "Markdown" in consigne
        assert "nb_slides" in consigne
        assert "récap" in consigne.lower()

    def test_xlsx_interdit_d_inventer_des_donnees(self, tmp_path):
        from app.services.skills.xlsx_generator import XlsxSkill

        consigne = XlsxSkill(output_dir=tmp_path).get_system_prompt_addition()
        assert "inventer" in consigne.lower()
        assert "récap" in consigne.lower()

    def test_le_bac_a_sable_transmet_nb_slides(self, tmp_path):
        from app.services.skills.code_executor import _build_namespace

        espace = _build_namespace(str(tmp_path / "x.pptx"), "Titre", "pptx", 7)
        assert espace["nb_slides"] == 7
        assert _build_namespace(str(tmp_path / "y.pptx"), "Titre", "pptx")["nb_slides"] == 10


# ---------------------------------------------------------------- version et migrations


class TestVersionEtMigrations:
    def test_la_version_est_coherente_et_en_semver(self):
        from app import __version__
        from app.config import settings

        assert __version__ == settings.app_version
        assert re.fullmatch(r"\d+\.\d+\.\d+(-[0-9A-Za-z.]+)?", __version__), __version__

    @pytest.mark.parametrize("colonne", ["provider", "extra_data"])
    def test_une_base_ancienne_recoit_la_colonne_manquante_des_messages(self, tmp_path, monkeypatch, colonne):
        from app.models import database as db_module
        from sqlalchemy import create_engine
        from sqlmodel import SQLModel

        monkeypatch.setattr(db_module, "db_encryption_enabled", lambda: False)
        chemin = tmp_path / "therese.db"
        moteur = create_engine(f"sqlite:///{chemin}")
        SQLModel.metadata.create_all(moteur)
        moteur.dispose()
        with closing(sqlite3.connect(chemin)) as conn:
            conn.execute(f"ALTER TABLE messages DROP COLUMN {colonne}")
            conn.commit()

        db_module.apply_adhoc_migrations(chemin)

        with closing(sqlite3.connect(chemin)) as conn:
            colonnes = {row[1] for row in conn.execute("PRAGMA table_info(messages)").fetchall()}
        assert colonne in colonnes, f"la migration ad hoc doit poser messages.{colonne}"


# ---------------------------------------------------------------- outil agenda et messages


class TestOutilAgenda:
    @pytest.mark.asyncio
    async def test_la_fenetre_par_defaut_couvre_trente_jours(self, db_session, monkeypatch):
        from app.services import workspace_tools

        recu: dict = {}

        class Fournisseur:
            async def list_events(self, calendar_id, time_min=None, time_max=None, max_results=50, **kwargs):
                recu.update(time_min=time_min, time_max=time_max)
                return [], None

        async def faux(session, calendar_id=None):
            return Fournisseur(), "cal", None

        monkeypatch.setattr(workspace_tools, "_get_calendar_provider", faux)
        await workspace_tools._list_calendar_events({}, db_session)
        assert recu, "le fournisseur doit être interrogé"
        assert (recu["time_max"] - recu["time_min"]).days >= 29, "les échéances à un mois doivent être visibles (P0-PROD-3)"

    @pytest.mark.asyncio
    async def test_sans_calendrier_la_consigne_interdit_d_inventer(self, db_session, monkeypatch):
        from app.services import workspace_tools

        async def aucun(session, calendar_id=None):
            return None, None, "aucun compte connecté"

        monkeypatch.setattr(workspace_tools, "_get_calendar_provider", aucun)
        texte = await workspace_tools._list_calendar_events({}, db_session)
        assert "AUCUN CALENDRIER" in texte and "invente" in texte.lower()


class TestOrdreDesMessages:
    @pytest.mark.asyncio
    async def test_a_date_egale_l_ordre_est_stable_par_identifiant(self, client, db_session):
        from app.models.entities import Conversation, Message

        conv = Conversation(id="conv-ordre", title="Ordre")
        db_session.add(conv)
        await db_session.flush()
        meme_instant = datetime(2026, 9, 7, 10, 0, tzinfo=UTC).replace(tzinfo=None)
        for identifiant in ("msg-c", "msg-a", "msg-b"):
            db_session.add(Message(id=identifiant, conversation_id="conv-ordre", role="user", content=identifiant, created_at=meme_instant))
        await db_session.commit()

        reponse = await client.get("/api/chat/conversations/conv-ordre/messages")
        assert reponse.status_code == 200, reponse.text
        assert [m["id"] for m in reponse.json()] == ["msg-a", "msg-b", "msg-c"], (
            "à date égale, l'ordre doit être déterministe (BUG-031 : created_at puis id)"
        )


# ---------------------------------------------------------------- limitation de débit et documents


class TestLimitationDeDebit:
    def test_le_limiteur_est_installe_sur_l_application(self):
        from app.main import app

        assert getattr(app.state, "limiter", None) is not None, "le rate limiting doit être configuré"


class TestValidationDesDocumentsProduits:
    def test_un_document_squelettique_est_refuse_et_un_document_rempli_accepte(self, tmp_path):
        import docx
        from app.services.skills.code_executor import _validate_document_content

        vide = tmp_path / "vide.docx"
        d = docx.Document()
        d.add_paragraph("Titre")
        d.save(vide)
        assert _validate_document_content(str(vide), "docx") is False, "un squelette sans contenu doit déclencher la relance (BUG-043)"

        rempli = tmp_path / "rempli.docx"
        d = docx.Document()
        d.add_heading("Programme", 1)
        for i in range(8):
            d.add_paragraph(f"Séance {i + 1} : objectifs, exercices et livrables détaillés pour les participants.")
        d.add_table(rows=3, cols=3)
        d.save(rempli)
        assert _validate_document_content(str(rempli), "docx") is True
