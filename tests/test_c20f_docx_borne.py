"""B-1160 : les fusions XML ne doivent pas amplifier le travail d'extraction."""

import subprocess
import sys
from copy import deepcopy

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


def creer_docx(path, *, span=1, lignes=1, fusion=False):
    doc = Document()
    doc.add_paragraph("Avant")
    table = doc.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "Contenu"
    tc = table._tbl.tr_lst[0].tc_lst[0]
    tc.get_or_add_tcPr().get_or_add_gridSpan().val = span
    if fusion:
        merge = OxmlElement("w:vMerge")
        merge.set(qn("w:val"), "restart")
        tc.get_or_add_tcPr().append(merge)
    for _ in range(lignes - 1):
        ligne = deepcopy(table._tbl.tr_lst[0])
        if fusion:
            ligne.tc_lst[0].tcPr.vMerge.val = "continue"
        table._tbl.append(ligne)
    doc.add_paragraph("Après")
    doc.save(path)


def extraire_borne(path):
    if sys.platform != "linux":
        pytest.skip("Cette reproduction exige le plafond mémoire RLIMIT_AS de Linux")
    # Le plafond est posé AVANT l'import. Un timeout tue et récolte le fils.
    code = """
import resource
import sys
from pathlib import Path
resource.setrlimit(resource.RLIMIT_AS, (1024**3, 1024**3))
sys.path.insert(0, 'src/backend')
from app.services.file_parser import _extract_docx
print(_extract_docx(Path(sys.argv[1])))
"""
    try:
        resultat = subprocess.run(
            [sys.executable, "-B", "-c", code, str(path)],
            capture_output=True, text=True, timeout=5,
        )
    except subprocess.TimeoutExpired:
        pytest.fail("Extraction DOCX supérieure à 5 s sous plafond de 1 Go")
    assert resultat.returncode == 0, resultat.stderr
    return resultat.stdout.strip()


@pytest.mark.parametrize("span,lignes,fusion,tronque", [
    (2_000_000_000, 1, False, True),
    (1, 800, True, False),
    (1, 1001, True, True),
])
def test_docx_hostile_borne(tmp_path, span, lignes, fusion, tronque):
    path = tmp_path / "devis.docx"
    creer_docx(path, span=span, lignes=lignes, fusion=fusion)
    texte = extraire_borne(path)
    assert "Avant" in texte and "Après" in texte
    if tronque:
        assert "DOCX tronqué" in texte
    else:
        assert texte == "Avant\n\nContenu\n\nAprès"


@pytest.mark.asyncio
async def test_reprise_sync_docx_tronque_ne_repart_pas(client, tmp_path, monkeypatch):
    from unittest.mock import AsyncMock

    from app.models.entities_sync import EtatOperation, EtatPlan
    from app.services import indexation
    from app.services import project_sync_service as svc
    from app.services.file_parser import extract_text

    dossier = tmp_path / "client"
    dossier.mkdir()
    creer_docx(dossier / "devis.docx", span=129)
    qdrant = AsyncMock()
    monkeypatch.setattr(indexation, "get_qdrant_service", lambda: qdrant)
    textes = []

    def extraire(path):
        texte = extract_text(path)
        textes.append(texte)
        return texte

    monkeypatch.setattr(indexation, "extract_text", extraire)
    projet = (await client.post("/api/memory/projects", json={"name": "Reprise DOCX"})).json()["id"]
    await svc.definir_racine(projet, str(dossier))
    plan = await svc.preparer_plan(projet)
    await svc._marquer_plan(plan.id, EtatPlan.EN_COURS)
    assert await svc.reprendre_applies_orphelins() == 1
    assert (await svc.lire_plan(plan.id)).etat == EtatPlan.APPLIQUE
    assert all(op.etat == EtatOperation.FAIT for op in await svc.lire_operations(plan.id))
    assert len(textes) == 1 and "DOCX tronqué" in textes[0]
    assert await svc.reprendre_applies_orphelins() == 0


@pytest.mark.parametrize("span,lignes,fusion,tronque", [
    (3, 1, False, False), (7, 4, False, False),
    (31, 3, True, False), (127, 1, False, False),
    (128, 2, True, False), (129, 1, False, True),
    (10_000_000, 50, False, True), (2_147_483_647, 1, False, True),
    (1, 999, True, False), (1, 1002, True, True),
])
def test_dix_rejeux_apres_correctif(tmp_path, span, lignes, fusion, tronque):
    path = tmp_path / "rejeu.docx"
    creer_docx(path, span=span, lignes=lignes, fusion=fusion)
    texte = extraire_borne(path)
    assert texte.startswith("Avant\n\n") and texte.endswith("\n\nAprès")
    assert ("DOCX tronqué" in texte) == tronque
    if not tronque:
        attendu = ["Contenu"] * (1 if fusion else lignes)
        assert texte == "\n\n".join(["Avant", *attendu, "Après"])


def test_budget_cellules_partage_entre_tableaux(tmp_path):
    doc = Document()
    table = doc.add_table(rows=1000, cols=10)
    table.cell(0, 0).text = "Premier tableau"
    doc.add_table(rows=1, cols=1).cell(0, 0).text = "Hors budget"
    path = tmp_path / "cellules.docx"
    doc.save(path)
    texte = extraire_borne(path)
    assert "Premier tableau" in texte
    assert "Hors budget" not in texte
    assert "DOCX tronqué" in texte
