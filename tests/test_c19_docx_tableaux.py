"""B-1103 : les tableaux font partie du texte indexable."""

import pytest
from app.services.file_parser import _extract_docx
from docx import Document


def test_tableaux_ordonnes_imbriques_et_cellules_fusionnees(tmp_path):
    doc = Document()
    doc.add_paragraph("Avant la grille")
    table = doc.add_table(rows=2, cols=3)
    table.cell(0, 0).text = "Mimavel"
    table.cell(0, 1).text = "7500 euros"
    table.cell(0, 2).text = "Livraison"
    table.cell(1, 0).merge(table.cell(1, 1)).text = "Fusion horizontale"
    table.cell(0, 2).merge(table.cell(1, 2)).text = "Fusion verticale"
    cellule = table.cell(1, 0)
    cellule.add_paragraph("Avant détail")
    cellule.add_table(rows=1, cols=2).rows[0].cells[0].text = "Détail imbriqué"
    cellule.add_paragraph("Après détail")
    doc.add_paragraph("Après la grille")
    path = tmp_path / "grille.docx"
    doc.save(path)

    texte = _extract_docx(path)

    assert "Mimavel\t7500 euros\tFusion verticale" in texte
    for mot in ["Fusion horizontale", "Fusion verticale", "Détail imbriqué"]:
        assert texte.count(mot) == 1
    positions = [texte.index(mot) for mot in [
        "Avant la grille", "Mimavel", "Avant détail", "Détail imbriqué",
        "Après détail", "Après la grille",
    ]]
    assert positions == sorted(positions)


@pytest.mark.parametrize("produit,prix", [
    ("Acajou", "128 euros"), ("Érable", "219 euros"), ("Palissandre", "381 euros"),
    ("Bouleau", "74 euros"), ("Cèdre", "805 euros"), ("Orme", "96 euros"),
    ("Tilleul", "147 euros"), ("Douglas", "633 euros"), ("Sapin", "42 euros"),
    ("Robinier", "958 euros"),
])
def test_rejeu_grilles(tmp_path, produit, prix):
    doc = Document()
    doc.add_paragraph("Catalogue")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = produit
    table.cell(0, 1).text = prix
    doc.add_paragraph("Fin catalogue")
    path = tmp_path / "catalogue.docx"
    doc.save(path)
    texte = _extract_docx(path)
    assert f"{produit}\t{prix}" in texte
    assert texte.index("Catalogue") < texte.index(produit) < texte.index("Fin catalogue")


def entourer_controle(element, profondeur=1):
    from docx.oxml import OxmlElement

    for _ in range(profondeur):
        controle = OxmlElement("w:sdt")
        contenu = OxmlElement("w:sdtContent")
        controle.append(contenu)
        element.addprevious(controle)
        contenu.append(element)
        element = controle


@pytest.mark.parametrize("dans_cellule", [False, True])
@pytest.mark.parametrize("profondeur", [1, 2])
def test_controles_de_contenu_traverses_dans_ordre(tmp_path, dans_cellule, profondeur):
    doc = Document()
    parent = doc.add_table(rows=1, cols=1).cell(0, 0) if dans_cellule else doc
    parent.add_paragraph("Avant")
    intro = parent.add_paragraph("Début du contrôle")
    table = parent.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Chêne"
    table.cell(0, 1).text = "7500 euros"
    fin = parent.add_paragraph("Fin du contrôle")
    for element in (intro._p, table._tbl, fin._p):
        entourer_controle(element, profondeur)
    parent.add_paragraph("Après")
    path = tmp_path / "controle.docx"
    doc.save(path)

    assert _extract_docx(path) == "Avant\n\nDébut du contrôle\n\nChêne\t7500 euros\n\nFin du contrôle\n\nAprès"
