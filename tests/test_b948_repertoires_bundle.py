"""B-948 (cycle 11, 23/09/2026) : sur Linux .deb, la génération DOCX échoue.

python-docx et python-pptx ouvrent leurs gabarits par un chemin relatif à leur
module, par exemple `docx/parts/../templates/default-footer.xml`. PyInstaller
range les modules dans l'archive PYZ : `docx/parts/` n'existe pas sur disque,
et Linux refuse de résoudre `..` à travers un dossier absent.

Le hook `runtime_hook_templates.py` crée ces dossiers au démarrage. Sur le
paquet .deb, le bundle est installé sous `/usr/lib/THERESE/` : le hook ne peut
pas écrire, et BUG-052 (0.4.6) s'est contenté d'ignorer l'erreur en annonçant
que « le packaging » créerait les dossiers. Rien ne le faisait : le .deb publié
de la 0.74.0 contient `docx/templates/default-footer.xml` mais pas
`docx/parts`, et `docx_generator._add_footer` passe par ce chemin à chaque
document. Preuves : `.app-loop/cycles/11/reproduce/`.

Correctif : le spec embarque un fichier témoin dans chacun de ces dossiers,
qui existent donc dans tout bundle, en lecture seule compris.
"""

from __future__ import annotations

import ast
import importlib.util
import os
import stat
import subprocess
import sys
from pathlib import Path

import pytest

RACINE = Path(__file__).resolve().parent.parent
BACKEND = RACINE / "src" / "backend"
SPEC = BACKEND / "backend.spec"
HOOK = BACKEND / "runtime_hook_templates.py"
MODULE = BACKEND / "repertoires_bundle.py"


def _module_repertoires():
    spec = importlib.util.spec_from_file_location("repertoires_bundle_b948", MODULE)
    assert spec and spec.loader, f"{MODULE} introuvable"
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def _repertoires_du_hook() -> tuple[str, ...]:
    for noeud in ast.walk(ast.parse(HOOK.read_text(encoding="utf-8"))):
        if isinstance(noeud, ast.For) and isinstance(noeud.iter, ast.Tuple):
            return tuple(e.value for e in noeud.iter.elts if isinstance(e, ast.Constant))
    pytest.fail("boucle des répertoires introuvable dans le hook")


class TestB948Structure:
    def test_le_spec_ajoute_les_temoins_aux_datas(self):
        arbre = ast.parse(SPEC.read_text(encoding="utf-8"))
        ajouts = [
            n for n in ast.walk(arbre)
            if isinstance(n, ast.AugAssign) and isinstance(n.target, ast.Name)
            and n.target.id == "datas" and "datas_repertoires_requis" in ast.unparse(n.value)
        ]
        assert ajouts, "backend.spec doit faire `datas += datas_repertoires_requis()` (B-948)"

    def test_le_module_et_le_hook_nomment_les_memes_repertoires(self):
        module = _module_repertoires()
        assert tuple(module.REPERTOIRES_REQUIS) == _repertoires_du_hook()
        assert {"docx/parts", "pptx/oxml", "pptx/shapes"} <= set(module.REPERTOIRES_REQUIS)

    def test_chaque_datas_pointe_un_temoin_reel_vers_un_repertoire_requis(self):
        module = _module_repertoires()
        datas = module.datas_repertoires_requis()
        assert [dest for _src, dest in datas] == list(module.REPERTOIRES_REQUIS)
        for src, _dest in datas:
            assert Path(src).is_file(), src

    def test_un_job_ci_linux_execute_la_contre_epreuve_reelle(self):
        """B-956 (revue Codex du diff) : sans job qui pose la variable, la preuve
        comportementale ne tournait nulle part en CI."""
        import re

        import yaml

        jobs = yaml.safe_load((RACINE / ".github" / "workflows" / "ci.yml").read_text(encoding="utf-8"))["jobs"]
        candidats = []
        for nom, job in jobs.items():
            if "ubuntu" not in str(job.get("runs-on", "")):
                continue
            env_job = {str(k): str(v) for k, v in (job.get("env") or {}).items()}
            for etape in job.get("steps", []):
                commande = str(etape.get("run", ""))
                env = {**env_job, **{str(k): str(v) for k, v in (etape.get("env") or {}).items()}}
                if "test_b948_repertoires_bundle.py" in commande and (
                    env.get("THERESE_TESTS_PYINSTALLER") == "1" or "THERESE_TESTS_PYINSTALLER=1" in commande
                ):
                    borne = re.search(r"--timeout=(\d+)", commande)
                    candidats.append((nom, int(borne.group(1)) if borne else None))
        assert candidats, "aucun job Linux ne lance la contre-épreuve avec THERESE_TESTS_PYINSTALLER=1"
        for nom, borne in candidats:
            assert borne is not None and borne >= 300, f"{nom} : --timeout trop court pour deux constructions PyInstaller"


# ------------------------------------------------------------ comportement réel

PROGRAMME = '''
import json, sys
res = {}
try:
    from docx import Document
    d = Document()
    f = d.sections[-1].footer
    f.is_linked_to_previous = False
    f.paragraphs[0].add_run("Généré par THÉRÈSE")
    res["docx_footer"] = "ok"
except Exception as exc:
    res["docx_footer"] = f"{type(exc).__name__}: {exc}"
print(json.dumps(res, ensure_ascii=False))
'''


def _construire_et_lancer_en_lecture_seule(tmp: Path, nom: str, datas: list[tuple[str, str]]) -> dict:
    import json

    tmp.mkdir(parents=True, exist_ok=True)
    source = tmp / f"{nom}.py"
    source.write_text(PROGRAMME, encoding="utf-8")
    commande = [
        sys.executable, "-m", "PyInstaller", "--onedir", "--noconfirm", "--clean",
        "--name", nom, "--collect-data", "docx", "--runtime-hook", str(HOOK),
        "--distpath", str(tmp / "dist"), "--workpath", str(tmp / "build"), "--specpath", str(tmp),
    ]
    for src, dest in datas:
        commande += ["--add-data", f"{src}{os.pathsep}{dest}"]
    construit = subprocess.run([*commande, str(source)], capture_output=True, text=True, timeout=600)
    assert construit.returncode == 0, construit.stderr[-2000:]
    bundle = tmp / "dist" / nom
    # Situation du .deb : installé sous /usr/lib, non inscriptible par l'utilisateur.
    for chemin in [bundle, *bundle.rglob("*")]:
        if not chemin.is_symlink():
            chemin.chmod(chemin.stat().st_mode & ~(stat.S_IWUSR | stat.S_IWGRP | stat.S_IWOTH))
    try:
        sortie = subprocess.run([str(bundle / nom)], capture_output=True, text=True, timeout=120)
    finally:
        for chemin in [bundle, *bundle.rglob("*")]:
            if not chemin.is_symlink():
                chemin.chmod(chemin.stat().st_mode | stat.S_IWUSR)
    assert sortie.returncode == 0, sortie.stderr[-2000:]
    return json.loads(sortie.stdout.strip().splitlines()[-1])


@pytest.mark.skipif(
    os.environ.get("THERESE_TESTS_PYINSTALLER") != "1" or not sys.platform.startswith("linux"),
    reason="Construit deux bundles PyInstaller (environ une minute) : poser THERESE_TESTS_PYINSTALLER=1 sous Linux",
)
def test_un_bundle_en_lecture_seule_genere_un_docx_avec_pied_de_page(tmp_path: Path):
    if importlib.util.find_spec("PyInstaller") is None:
        pytest.skip("PyInstaller absent")
    module = _module_repertoires()

    # Contre-épreuve : sans les témoins, l'échec du .deb 0.74 se reproduit.
    sans = _construire_et_lancer_en_lecture_seule(tmp_path / "sans", "sans_temoins", [])
    assert sans["docx_footer"].startswith("FileNotFoundError"), sans

    avec = _construire_et_lancer_en_lecture_seule(
        tmp_path / "avec", "avec_temoins", module.datas_repertoires_requis()
    )
    assert avec["docx_footer"] == "ok", avec
