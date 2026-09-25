"""B-1453 (recette P-146, lot 3) : la page de titre d'une présentation
produite par THÉRÈSE portait le sous-titre « Synoptia - L'entrepreneur
augmenté », le slogan de l'éditeur, sur le document de l'utilisatrice
(« Offre d'accompagnement numérique d'Atelier Ménard »). La page de titre
ne porte que le titre demandé.

Le pied « Généré par THERESE - Synoptia » de la dernière page est une autre
question (choix de marque), posée à Ludo."""

from pathlib import Path

import pytest
from app.services.skills.base import SkillParams
from pptx import Presentation


@pytest.mark.asyncio
async def test_la_page_de_titre_ne_porte_pas_le_slogan_de_l_editeur(tmp_path: Path):
    from app.services.skills.pptx_generator import PptxSkill

    skill = PptxSkill(output_dir=tmp_path)
    markdown = "## Contexte\n- Atelier de menuiserie\n- Trois salariés\n\n---\n\n## Offre\n- Agenda partagé\n"
    resultat = await skill.execute(SkillParams(title="Offre Atelier Ménard", content=markdown))
    page_de_titre = Presentation(str(resultat.file_path)).slides[0]
    textes = [forme.text_frame.text for forme in page_de_titre.shapes if forme.has_text_frame]
    assert "Offre Atelier Ménard" in textes
    assert not any("entrepreneur augmenté" in texte or "Synoptia" in texte for texte in textes), textes
