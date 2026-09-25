"""B-1451 (recette P-146, lot 3, KO-4) : la puce « Document Word » ne
produisait rien. `docx.Document()` ouvre toujours le modèle `default.docx`
rangé dans python-docx ; le bac à sable borne `open()` au dossier de sortie
(passe 4), la lecture était refusée (PackageNotFoundError), et toute
génération Word par exécution de code échouait, application installée
comprise.

Les modèles livrés AVEC les bibliothèques Office autorisées (docx, pptx,
openpyxl) se lisent ; rien d'autre, et jamais en écriture."""

import queue
from pathlib import Path

from app.services.skills.code_executor import _run_generation_in_subprocess

DOCX = (
    "doc = Document()\n"
    "doc.add_heading(title, level=0)\n"
    "doc.add_paragraph('Relances clients : un point chaque lundi.')\n"
    "doc.add_paragraph('Chaque relance a une date.')\n"
    "doc.add_paragraph('Les relances faites sortent du brief.')\n"
    "doc.save(output_path)\n"
)


def _executer(code: str, sortie: Path, fmt: str) -> tuple[str, str]:
    rq: queue.Queue = queue.Queue()
    _run_generation_in_subprocess(code, str(sortie), "Note de cadrage", fmt, 10, rq)
    return rq.get_nowait()


def test_un_document_word_se_genere(tmp_path):
    sortie = tmp_path / "note.docx"
    statut, detail = _executer(DOCX, sortie, "docx")
    assert statut == "ok", detail
    assert sortie.exists() and sortie.stat().st_size > 0


def test_le_garde_refuse_toujours_une_lecture_hors_des_bibliotheques(tmp_path):
    import builtins

    import pytest
    from app.services.skills.code_executor import _installer_garde_fs

    secret = tmp_path.parent / "secret-b1451.txt"
    secret.write_text("ne pas lire")
    restaurer = _installer_garde_fs(str(tmp_path / "sortie" / "x.docx"))
    try:
        with pytest.raises(PermissionError), builtins.open(secret):
            pass
    finally:
        restaurer()


def test_le_garde_lit_le_modele_mais_ne_l_ecrit_jamais(tmp_path):
    import builtins

    import docx
    import pytest
    from app.services.skills.code_executor import _installer_garde_fs

    modele = Path(docx.__file__).parent / "templates" / "default.docx"
    restaurer = _installer_garde_fs(str(tmp_path / "x.docx"))
    try:
        with builtins.open(modele, "rb") as flux:
            assert flux.read(2) == b"PK"
        with pytest.raises(PermissionError), builtins.open(modele, "ab"):
            pass
    finally:
        restaurer()
