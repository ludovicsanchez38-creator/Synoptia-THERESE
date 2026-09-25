"""
THERESE v2 - CRM Import Service

Multi-format import for CRM data (CSV, Excel, JSON).
Part of the "Local First" architecture.
"""

import csv
import io
import json
import logging
import math
from dataclasses import dataclass, field
from datetime import UTC, datetime
from typing import Any, Literal

from app.models.entities import Contact, Deliverable, Project, generate_uuid
from app.models.schemas import adresse_unique_valide, perimetre_normalise
from app.services.crm_utils import (
    DELIVERABLE_STATUS_MAP,
    ETAPES_PIPELINE,
    PROJECT_STATUS_MAP,
    cle_de_statut,
    etiquettes_ecartees,
    etiquettes_lues,
)
from app.services.formules_tableur import neutraliser_formule
from openpyxl import load_workbook
from sqlalchemy.ext.asyncio import AsyncSession
from sqlmodel import select

logger = logging.getLogger(__name__)


ImportFormat = Literal["csv", "xlsx", "json"]


@dataclass
class ImportError:
    """Error encountered during import."""
    row: int
    column: str | None
    message: str
    data: dict | None = None


@dataclass
class ImportResult:
    """Result of an import operation."""
    success: bool
    created: int = 0
    updated: int = 0
    skipped: int = 0
    errors: list[ImportError] = field(default_factory=list)
    total_rows: int = 0

    @property
    def message(self) -> str:
        """Human-readable summary."""
        # B-1278 : un signalement (cellule écartée) ne fait pas oublier ce qui
        # a été enregistré ; le message compte toujours les lignes.
        bilan = f"{self.created} créés, {self.updated} mis à jour, {self.skipped} ignorés"
        # B-1289 : rien lu (fichier illisible) n'est pas un import terminé.
        if not self.success and self.total_rows == 0 and self.errors:
            return f"Import impossible : {self.errors[0].message}"
        if self.success:
            return f"Import terminé : {bilan}"
        return f"Import terminé avec {len(self.errors)} signalement(s) : {bilan}, sur {self.total_rows} lignes"


@dataclass
class ImportPreview:
    """Preview of import before execution."""
    total_rows: int
    sample_rows: list[dict]
    detected_columns: list[str]
    column_mapping: dict[str, str]
    validation_errors: list[ImportError]
    can_import: bool


# ============================================================
# Column Mappings
# ============================================================

# Flexible column name mappings (source -> internal)
CONTACT_COLUMN_MAPPING = {
    # ID
    "id": "id",
    "ID": "id",
    "identifiant": "id",
    # First name
    "first_name": "first_name",
    "prenom": "first_name",
    "Prenom": "first_name",
    "prénom": "first_name",
    "Prénom": "first_name",
    "firstname": "first_name",
    "FirstName": "first_name",
    # Last name
    "last_name": "last_name",
    "nom": "last_name",
    "Nom": "last_name",
    "lastname": "last_name",
    "LastName": "last_name",
    # Company
    "company": "company",
    "entreprise": "company",
    "Entreprise": "company",
    "societe": "company",
    "Societe": "company",
    "société": "company",
    "Société": "company",
    # Email
    "email": "email",
    "Email": "email",
    "e-mail": "email",
    "E-mail": "email",
    "courriel": "email",
    # Phone
    "phone": "phone",
    "telephone": "phone",
    "Telephone": "phone",
    "téléphone": "phone",
    "Téléphone": "phone",
    "tel": "phone",
    "Tel": "phone",
    # Address
    "address": "address",
    "adresse": "address",
    "Adresse": "address",
    # Stage
    "stage": "stage",
    "Stage": "stage",
    "etape": "stage",
    "Etape": "stage",
    # Score
    "score": "score",
    "Score": "score",
    # Source
    "source": "source",
    "Source": "source",
    "origine": "source",
    # Tags
    "tags": "tags",
    "Tags": "tags",
    "etiquettes": "tags",
    # Notes
    "notes": "notes",
    "Notes": "notes",
    "commentaires": "notes",
    # Long-term state exported by THERESE
    "extra_data": "extra_data",
    "Donnees supplementaires": "extra_data",
    "last_interaction": "last_interaction",
    "Derniere interaction": "last_interaction",
    "next_follow_up": "next_follow_up",
    "Prochaine relance": "next_follow_up",
    "rgpd_base_legale": "rgpd_base_legale",
    "Base legale RGPD": "rgpd_base_legale",
    "rgpd_date_collecte": "rgpd_date_collecte",
    "Date collecte RGPD": "rgpd_date_collecte",
    "rgpd_date_expiration": "rgpd_date_expiration",
    "Date expiration RGPD": "rgpd_date_expiration",
    "rgpd_consentement": "rgpd_consentement",
    "Consentement RGPD": "rgpd_consentement",
    "purge_excluded": "purge_excluded",
    "Exclu purge RGPD": "purge_excluded",
    "scope": "scope",
    "Perimetre": "scope",
    "scope_id": "scope_id",
    "ID perimetre": "scope_id",
    "created_at": "created_at",
    "Date creation": "created_at",
    "updated_at": "updated_at",
    "Date modification": "updated_at",
}

PROJECT_COLUMN_MAPPING = {
    # ID
    "id": "id",
    "ID": "id",
    # Name
    "name": "name",
    "Name": "name",
    "nom": "name",
    "Nom": "name",
    "titre": "name",
    "Titre": "name",
    # Description
    "description": "description",
    "Description": "description",
    # Contact ID
    "contact_id": "contact_id",
    "ContactID": "contact_id",
    "ClientID": "contact_id",
    "client_id": "contact_id",
    "ID Contact": "contact_id",
    # Status
    "status": "status",
    "Status": "status",
    "statut": "status",
    "Statut": "status",
    # Budget
    "budget": "budget",
    "Budget": "budget",
    # Notes
    "notes": "notes",
    "Notes": "notes",
    # Tags
    "tags": "tags",
    "Tags": "tags",
}

DELIVERABLE_COLUMN_MAPPING = {
    # ID
    "id": "id",
    "ID": "id",
    # Project ID
    "project_id": "project_id",
    "ProjectID": "project_id",
    "ID Projet": "project_id",
    # Title
    "title": "title",
    "Title": "title",
    "titre": "title",
    "Titre": "title",
    "nom": "title",
    "Nom": "title",
    # Description
    "description": "description",
    "Description": "description",
    # Status
    "status": "status",
    "Status": "status",
    "statut": "status",
    "Statut": "status",
    # Due date
    "due_date": "due_date",
    "DueDate": "due_date",
    "Date echeance": "due_date",
    "echeance": "due_date",
}


# ============================================================
# Parsing Helpers
# ============================================================


def _detect_format(content: bytes, filename: str | None = None) -> ImportFormat:
    """Detect file format from content or filename."""
    if filename:
        if filename.endswith(".csv"):
            return "csv"
        elif filename.endswith(".xlsx") or filename.endswith(".xls"):
            return "xlsx"
        elif filename.endswith(".json"):
            return "json"

    # Try to detect from content
    try:
        content.decode("utf-8")
        # Check if JSON
        stripped = content.strip()
        if stripped.startswith(b"[") or stripped.startswith(b"{"):
            return "json"
        return "csv"
    except UnicodeDecodeError:
        return "xlsx"


def _parse_csv(content: bytes) -> list[dict]:
    """Parse CSV content to list of dicts."""
    # Try different encodings
    for encoding in ["utf-8-sig", "utf-8", "latin-1", "cp1252"]:
        try:
            text = content.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    else:
        raise ValueError("Impossible de decoder le fichier CSV")

    reader = csv.DictReader(io.StringIO(text))
    # B-551 (05/09/2026) : une ligne plus longue que l'en-tête range ses
    # valeurs excédentaires sous la clé None (comportement de DictReader).
    # Cette clé cassait l'aperçu d'import (detected_columns exige des chaînes)
    # alors que l'import réel l'ignorait : les deux digèrent le même fichier.
    return [
        {cle: valeur for cle, valeur in ligne.items() if cle is not None}
        for ligne in reader
    ]


def _parse_xlsx(content: bytes, sheet_name: str | None = None) -> list[dict]:
    """Parse Excel content to list of dicts."""
    wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)

    if sheet_name and sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
    else:
        ws = wb.active

    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []

    headers = [str(h).strip() if h else f"col_{i}" for i, h in enumerate(rows[0])]
    data = []

    for row in rows[1:]:
        if all(cell is None for cell in row):
            continue
        row_dict = {}
        for i, cell in enumerate(row):
            if i < len(headers):
                row_dict[headers[i]] = cell
        data.append(row_dict)

    return data


def _parse_json(content: bytes) -> list[dict]:
    """Parse JSON content to list of dicts."""
    data = json.loads(content.decode("utf-8"))
    if isinstance(data, dict):
        # Check for nested structure
        if "contacts" in data:
            data = data["contacts"]
        elif "projects" in data:
            data = data["projects"]
        elif "deliverables" in data:
            data = data["deliverables"]
        else:
            return [data]
    # B-1301 : un nombre, un texte ou null tombait plus loin en TypeError (500).
    if not isinstance(data, list):
        raise ValueError("Le fichier JSON doit contenir une liste de fiches (ou un objet).")
    # B-1312 : [1, 2] passait, et l'aperçu tombait sur `raw_data[0].keys()`.
    if any(not isinstance(fiche, dict) for fiche in data):
        raise ValueError("Chaque fiche du fichier JSON doit être un objet.")
    return data


# Field length limits (SEC-017)
FIELD_MAX_LENGTHS: dict[str, int] = {
    "first_name": 200,
    "last_name": 200,
    "company": 300,
    "email": 320,
    "phone": 50,
    "address": 1000,
    "stage": 50,
    "source": 200,
    "tags": 1000,
    "notes": 5000,
    "extra_data": 10000,
    "rgpd_base_legale": 100,
    "scope": 50,
    "scope_id": 200,
    "name": 500,
    "title": 500,
    "description": 5000,
    "status": 50,
}



def _sanitize_field(value: Any, field_name: str | None = None) -> Any:
    """
    Sanitize a field value for safe storage (SEC-017).

    - Strip whitespace
    - Enforce length limits
    - Neutralize formula injection prefixes (CSV injection defense)
    - Remove null bytes
    """
    if value is None:
        return None
    if not isinstance(value, str):
        return value
    # Strip and remove null bytes
    value = value.strip().replace("\x00", "")
    if not value:
        return None
    # SEC-017, B-1120 : même règle qu'à l'export ; un téléphone « +33… » ou un
    # budget « -500 » restent intacts, « +1+cmd|… » est désamorcé.
    value = neutraliser_formule(value)
    # Enforce length limit
    if field_name and field_name in FIELD_MAX_LENGTHS:
        max_len = FIELD_MAX_LENGTHS[field_name]
        if len(value) > max_len:
            value = value[:max_len]
    return value


def _perimetre_ou_none(valeur: Any) -> str | None:
    """B-1087, B-1165 : périmètre normalisé, ou None s'il est vide ou inconnu
    (un import ne tombe pas pour une valeur de périmètre)."""
    if not isinstance(valeur, str) or not valeur.strip():
        return None
    try:
        # Variable annotée : mypy lit les imports `app.*` comme Any.
        perimetre: str | None = perimetre_normalise(valeur)
    except ValueError:
        return None
    return perimetre


def _map_columns(row: dict, mapping: dict[str, str]) -> dict[str, Any]:
    """Map source columns to internal column names and sanitize values (SEC-017)."""
    result = {}
    for source_col, value in row.items():
        internal_col = mapping.get(source_col, source_col)
        if internal_col in mapping.values():
            result[internal_col] = _sanitize_field(value, internal_col)
    return result


def _parse_value(value: Any, field_type: str) -> Any:
    """Parse and validate a value based on field type."""
    if value is None or (isinstance(value, str) and not value.strip()):
        return None

    if field_type == "int":
        try:
            return int(float(str(value).strip()))
        except (ValueError, TypeError, OverflowError):
            # B-1238 : « inf » levait OverflowError.
            return None

    if field_type == "float":
        try:
            nombre = float(str(value).strip().replace(",", "."))
        except (ValueError, TypeError):
            return None
        # B-1238 : un budget infini ou NaN n'est pas un nombre (B-1219).
        return nombre if math.isfinite(nombre) else None

    if field_type == "datetime":
        if isinstance(value, datetime):
            return value
        if isinstance(value, str):
            try:
                return datetime.fromisoformat(value.strip().replace("Z", "+00:00"))
            except ValueError:
                pass
            for fmt in [
                "%Y-%m-%dT%H:%M:%S.%fZ",
                "%Y-%m-%dT%H:%M:%SZ",
                "%Y-%m-%dT%H:%M:%S",
                "%Y-%m-%d",
                "%d/%m/%Y",
                "%d-%m-%Y",
            ]:
                try:
                    return datetime.strptime(value.strip(), fmt)
                except ValueError:
                    continue
        return None

    if field_type == "bool":
        if isinstance(value, bool):
            return value
        if isinstance(value, (int, float)):
            return bool(value)
        normalized = str(value).strip().lower()
        if normalized in {"1", "true", "oui", "yes", "vrai"}:
            return True
        if normalized in {"0", "false", "non", "no", "faux"}:
            return False
        return None

    if field_type == "json":
        if isinstance(value, (dict, list)):
            return json.dumps(value, ensure_ascii=False, separators=(",", ":"))
        return str(value).strip()

    if field_type == "tags":
        # B-1272 : même règle que l'import JSON des sauvegardes.
        tags = etiquettes_lues(value)
        return json.dumps(tags) if tags else None

    # String
    return str(value).strip() if value else None


def _lire_les_etiquettes(valeur: Any, ligne: int, mapped: dict[str, Any]) -> tuple[str | None, "ImportError | None"]:
    """B-1286 : étiquettes retenues (JSON) et signalement de celles écartées."""
    lues = etiquettes_lues(valeur)
    ecartees = etiquettes_ecartees(valeur)
    signalement = (
        ImportError(
            row=ligne,
            column="tags",
            message=f"{ecartees} étiquette(s) illisible(s), non enregistrée(s)",
            data=mapped,
        )
        if ecartees
        else None
    )
    return (json.dumps(lues) if lues else None), signalement


def _validate_contact(data: dict) -> list[str]:
    """Validate contact data, return list of errors."""
    errors = []
    if not data.get("first_name") and not data.get("last_name") and not data.get("company"):
        errors.append("Au moins un nom ou une entreprise est requis")
    # B-1148 : une adresse double ou douteuse n'écarte plus la fiche ; elle est
    # retirée par _ecarter_adresse_douteuse, comme aux autres portes (B-1074).
    return errors


def _ecarter_adresse_douteuse(data: dict[str, Any]) -> str | None:
    """B-1074, B-1148 : une fiche porte une seule adresse de la forme nom@domaine.

    Une valeur douteuse (« jean@a.fr, compta@a.fr ») est retirée de la ligne,
    qui s'importe sans adresse ; sur une fiche existante, l'adresse valide déjà
    enregistrée n'est donc pas effacée (motif de B-1107). Rend la note à
    montrer, ou None."""
    valeur = data.get("email")
    if not valeur or adresse_unique_valide(str(valeur)):
        return None
    data.pop("email", None)
    return (
        f"Adresse e-mail ignorée (« {str(valeur)[:80]} ») : une seule adresse par fiche. "
        "La fiche est importée sans cette adresse."
    )


def _validate_project(data: dict) -> list[str]:
    """Validate project data, return list of errors."""
    errors = []
    if not data.get("name"):
        errors.append("Le nom du projet est requis")
    return errors


def _validate_deliverable(data: dict) -> list[str]:
    """Validate deliverable data, return list of errors."""
    errors = []
    if not data.get("title"):
        errors.append("Le titre du livrable est requis")
    if not data.get("project_id"):
        errors.append("L'ID du projet est requis")
    return errors


# ============================================================
# Import Service
# ============================================================


class CRMImportService:
    """
    Service for importing CRM data from multiple formats.

    Supports CSV, Excel (XLSX), and JSON imports for:
    - Contacts
    - Projects
    - Deliverables
    """

    def __init__(self, session: AsyncSession):
        """
        Initialize import service.

        Args:
            session: AsyncSession for database access
        """
        self.session = session

    async def preview_contacts(
        self,
        content: bytes,
        filename: str | None = None,
        custom_mapping: dict[str, str] | None = None,
    ) -> ImportPreview:
        """
        Preview contact import without executing.

        Args:
            content: File content as bytes
            filename: Original filename for format detection
            custom_mapping: Custom column mapping to override defaults

        Returns:
            ImportPreview with sample data and validation results
        """
        format_type = _detect_format(content, filename)
        mapping = {**CONTACT_COLUMN_MAPPING, **(custom_mapping or {})}

        try:
            if format_type == "csv":
                raw_data = _parse_csv(content)
            elif format_type == "xlsx":
                raw_data = _parse_xlsx(content)
            else:
                raw_data = _parse_json(content)
        except Exception as e:
            return ImportPreview(
                total_rows=0,
                sample_rows=[],
                detected_columns=[],
                column_mapping={},
                validation_errors=[ImportError(row=0, column=None, message=str(e))],
                can_import=False,
            )

        if not raw_data:
            return ImportPreview(
                total_rows=0,
                sample_rows=[],
                detected_columns=[],
                column_mapping={},
                validation_errors=[ImportError(row=0, column=None, message="Aucune donnee trouvee")],
                can_import=False,
            )

        detected_columns = list(raw_data[0].keys()) if raw_data else []
        used_mapping = {col: mapping.get(col, col) for col in detected_columns if col in mapping}

        validation_errors = []
        sample_rows = []

        bloquantes = []
        for idx, row in enumerate(raw_data[:5]):
            mapped = _map_columns(row, mapping)
            note = _ecarter_adresse_douteuse(mapped)
            sample_rows.append(mapped)
            if note:
                # B-1148 : prévenue, pas bloquante.
                validation_errors.append(ImportError(row=idx + 1, column="email", message=note, data=mapped))

            errors = _validate_contact(mapped)
            for error in errors:
                erreur = ImportError(row=idx + 1, column=None, message=error, data=mapped)
                validation_errors.append(erreur)
                bloquantes.append(erreur)

        can_import = len(bloquantes) == 0 or all(
            err.row > 5 for err in bloquantes
        )

        return ImportPreview(
            total_rows=len(raw_data),
            sample_rows=sample_rows,
            detected_columns=detected_columns,
            column_mapping=used_mapping,
            validation_errors=validation_errors[:10],
            can_import=can_import,
        )

    async def import_contacts(
        self,
        content: bytes,
        filename: str | None = None,
        custom_mapping: dict[str, str] | None = None,
        update_existing: bool = True,
    ) -> ImportResult:
        """
        Import contacts from file.

        Args:
            content: File content as bytes
            filename: Original filename for format detection
            custom_mapping: Custom column mapping to override defaults
            update_existing: Whether to update existing contacts by ID

        Returns:
            ImportResult with counts and errors
        """
        format_type = _detect_format(content, filename)
        mapping = {**CONTACT_COLUMN_MAPPING, **(custom_mapping or {})}

        try:
            if format_type == "csv":
                raw_data = _parse_csv(content)
            elif format_type == "xlsx":
                raw_data = _parse_xlsx(content)
            else:
                raw_data = _parse_json(content)
        except Exception as e:
            return ImportResult(
                success=False,
                errors=[ImportError(row=0, column=None, message=str(e))],
            )

        result = ImportResult(success=True, total_rows=len(raw_data))

        for idx, row in enumerate(raw_data):
            try:
                mapped = _map_columns(row, mapping)
                note = _ecarter_adresse_douteuse(mapped)
                if note:
                    result.errors.append(ImportError(row=idx + 1, column="email", message=note, data=mapped))

                # Validate
                errors = _validate_contact(mapped)
                if errors:
                    result.errors.append(ImportError(
                        row=idx + 1,
                        column=None,
                        message="; ".join(errors),
                        data=mapped,
                    ))
                    result.skipped += 1
                    continue

                # Check for existing contact
                contact_id = mapped.get("id")
                existing = None

                if contact_id:
                    stmt = select(Contact).where(Contact.id == contact_id)
                    db_result = await self.session.execute(stmt)
                    existing = db_result.scalar_one_or_none()

                # B-1262 : même règle que le tableur (B-1187) : seule une étape
                # du pipeline est retenue ; une autre rendait la fiche
                # invisible des colonnes. Elle ne remplace rien et figure au
                # rapport.
                cellule_etape = str(mapped.get("stage") or "").strip()
                etape = cellule_etape.lower() if cellule_etape.lower() in ETAPES_PIPELINE else None
                etape_ecartee = (
                    ImportError(
                        row=idx + 1,
                        column="stage",
                        message=f"Étape « {cellule_etape} » inconnue du pipeline, non enregistrée",
                        data=mapped,
                    )
                    if cellule_etape and etape is None
                    else None
                )

                if existing and update_existing:
                    # Un champ présent dans l'export fait foi, y compris une
                    # valeur vide ou un score nul : sinon l'aller-retour
                    # invente silencieusement un autre état.
                    for field_name in (
                        "first_name",
                        "last_name",
                        "company",
                        "email",
                        "phone",
                        "address",
                        "source",
                        "notes",
                        "rgpd_base_legale",
                        "scope_id",
                    ):
                        if field_name in mapped:
                            setattr(existing, field_name, mapped[field_name])
                    # B-1087 : étape et périmètre sont NOT NULL ; une cellule
                    # vide écrivait None et le commit unique faisait échouer
                    # tout l'import. Vide ou inconnu ne remplace rien ; le
                    # périmètre suit la règle de B-1165.
                    if etape:
                        existing.stage = etape
                    if etape_ecartee:
                        result.errors.append(etape_ecartee)
                    perimetre = _perimetre_ou_none(mapped.get("scope"))
                    if perimetre:
                        existing.scope = perimetre
                    if "score" in mapped:
                        score = _parse_value(mapped["score"], "int")
                        if score is not None:
                            existing.score = score
                    if "tags" in mapped:
                        # B-1286 : toutes écartées, elles ne remplacent rien ;
                        # une cellule vide garde son sens de miroir.
                        tags_lues, signalement_tags = _lire_les_etiquettes(mapped["tags"], idx + 1, mapped)
                        if tags_lues is not None or signalement_tags is None:
                            existing.tags = tags_lues
                        if signalement_tags:
                            result.errors.append(signalement_tags)
                    if "extra_data" in mapped:
                        existing.extra_data = _parse_value(mapped["extra_data"], "json")
                    for field_name in (
                        "last_interaction",
                        "next_follow_up",
                        "rgpd_date_collecte",
                        "rgpd_date_expiration",
                    ):
                        if field_name in mapped:
                            # B-1302 : vide = miroir (efface) ; illisible = ne
                            # remplace rien, et se dit (rgpd_date_expiration
                            # est la date de purge RGPD de la fiche).
                            cellule_date = str(mapped[field_name] or "").strip()
                            date_lue = _parse_value(mapped[field_name], "datetime")
                            if date_lue is None and cellule_date:
                                result.errors.append(ImportError(
                                    row=idx + 1,
                                    column=field_name,
                                    message=f"Date « {cellule_date} » illisible, non enregistrée",
                                    data=mapped,
                                ))
                            else:
                                setattr(existing, field_name, date_lue)
                    # B-1121 : date de création NOT NULL ; vide ou illisible,
                    # elle ne remplace pas celle de la fiche.
                    cree_le = _parse_value(mapped.get("created_at"), "datetime")
                    if cree_le is not None:
                        existing.created_at = cree_le
                    for field_name in ("rgpd_consentement", "purge_excluded"):
                        if field_name in mapped:
                            parsed_bool = _parse_value(mapped[field_name], "bool")
                            if parsed_bool is not None:
                                setattr(existing, field_name, parsed_bool)

                    existing.updated_at = (
                        _parse_value(mapped.get("updated_at"), "datetime") or datetime.now(UTC)
                    )
                    self.session.add(existing)
                    result.updated += 1

                elif existing and not update_existing:
                    result.skipped += 1

                else:
                    # Create new
                    score = _parse_value(mapped.get("score"), "int")
                    tags_de_la_ligne, signalement_tags = _lire_les_etiquettes(mapped.get("tags"), idx + 1, mapped)
                    # B-1313 : à la création aussi, une date illisible se dit.
                    dates_ecartees = [
                        ImportError(
                            row=idx + 1,
                            column=champ_date,
                            message=f"Date « {str(mapped[champ_date]).strip()} » illisible, non enregistrée",
                            data=mapped,
                        )
                        for champ_date in ("last_interaction", "next_follow_up", "rgpd_date_collecte", "rgpd_date_expiration")
                        if str(mapped.get(champ_date) or "").strip()
                        and _parse_value(mapped[champ_date], "datetime") is None
                    ]
                    created_at = _parse_value(mapped.get("created_at"), "datetime")
                    updated_at = _parse_value(mapped.get("updated_at"), "datetime")
                    contact = Contact(
                        id=contact_id or generate_uuid(),
                        first_name=mapped.get("first_name"),
                        last_name=mapped.get("last_name"),
                        company=mapped.get("company"),
                        email=mapped.get("email"),
                        phone=mapped.get("phone"),
                        address=mapped.get("address"),
                        stage=etape or "contact",
                        score=score if score is not None else 50,
                        source=mapped.get("source"),
                        tags=tags_de_la_ligne,
                        notes=mapped.get("notes"),
                        extra_data=_parse_value(mapped.get("extra_data"), "json"),
                        last_interaction=_parse_value(
                            mapped.get("last_interaction"), "datetime"
                        ),
                        next_follow_up=_parse_value(
                            mapped.get("next_follow_up"), "datetime"
                        ),
                        rgpd_base_legale=mapped.get("rgpd_base_legale"),
                        rgpd_date_collecte=_parse_value(
                            mapped.get("rgpd_date_collecte"), "datetime"
                        ),
                        rgpd_date_expiration=_parse_value(
                            mapped.get("rgpd_date_expiration"), "datetime"
                        ),
                        rgpd_consentement=(
                            _parse_value(mapped.get("rgpd_consentement"), "bool")
                            or False
                        ),
                        purge_excluded=(
                            _parse_value(mapped.get("purge_excluded"), "bool")
                            or False
                        ),
                        scope=_perimetre_ou_none(mapped.get("scope")) or "global",
                        scope_id=mapped.get("scope_id"),
                        created_at=created_at or datetime.now(UTC),
                        updated_at=updated_at or datetime.now(UTC),
                    )
                    self.session.add(contact)
                    result.created += 1
                    if etape_ecartee:
                        result.errors.append(etape_ecartee)
                    if signalement_tags:
                        result.errors.append(signalement_tags)
                    result.errors.extend(dates_ecartees)

            except Exception as e:
                logger.error(f"Error importing contact row {idx + 1}: {e}")
                result.errors.append(ImportError(row=idx + 1, column=None, message=str(e)))
                result.skipped += 1

        await self.session.commit()

        result.success = len(result.errors) == 0
        logger.info(f"Contact import: {result.message}")

        return result

    async def import_projects(
        self,
        content: bytes,
        filename: str | None = None,
        custom_mapping: dict[str, str] | None = None,
        update_existing: bool = True,
    ) -> ImportResult:
        """
        Import projects from file.

        Args:
            content: File content as bytes
            filename: Original filename for format detection
            custom_mapping: Custom column mapping to override defaults
            update_existing: Whether to update existing projects by ID

        Returns:
            ImportResult with counts and errors
        """
        format_type = _detect_format(content, filename)
        mapping = {**PROJECT_COLUMN_MAPPING, **(custom_mapping or {})}

        try:
            if format_type == "csv":
                raw_data = _parse_csv(content)
            elif format_type == "xlsx":
                raw_data = _parse_xlsx(content, sheet_name="Projets")
            else:
                raw_data = _parse_json(content)
        except Exception as e:
            return ImportResult(
                success=False,
                errors=[ImportError(row=0, column=None, message=str(e))],
            )

        result = ImportResult(success=True, total_rows=len(raw_data))

        # Status mapping

        for idx, row in enumerate(raw_data):
            try:
                mapped = _map_columns(row, mapping)

                errors = _validate_project(mapped)
                if errors:
                    result.errors.append(ImportError(
                        row=idx + 1,
                        column=None,
                        message="; ".join(errors),
                        data=mapped,
                    ))
                    result.skipped += 1
                    continue

                project_id = mapped.get("id")
                existing = None

                if project_id:
                    stmt = select(Project).where(Project.id == project_id)
                    db_result = await self.session.execute(stmt)
                    existing = db_result.scalar_one_or_none()

                # Validate contact_id exists
                contact_id = mapped.get("contact_id")
                if contact_id:
                    stmt = select(Contact).where(Contact.id == contact_id)
                    db_result = await self.session.execute(stmt)
                    if not db_result.scalar_one_or_none():
                        contact_id = None  # Don't link to non-existent contact

                # Parse status
                raw_status = str(mapped.get("status") or "").lower().strip()
                # B-1315 : même table et même clé que la synchro tableur.
                statut_reconnu = PROJECT_STATUS_MAP.get(cle_de_statut(raw_status))
                status = statut_reconnu or "active"
                # B-1307 : inconnu, le statut ne remplace rien (B-1083, B-1106)
                # et prend le défaut à la création ; il se dit au rapport.
                if raw_status and not statut_reconnu and not (existing and not update_existing):
                    result.errors.append(ImportError(
                        row=idx + 1,
                        column="status",
                        message=f"Statut « {mapped.get('status')} » inconnu, non enregistré",
                        data=mapped,
                    ))

                # B-1256 : une cellule de budget illisible (« inf », « NaN »,
                # texte) ne remplace rien et figure au rapport ; elle effaçait
                # en silence le budget existant.
                # B-1274 : `or ""` faisait d'un 0 numérique une cellule vide.
                brut_budget = mapped.get("budget")
                cellule_budget = "" if brut_budget is None else str(brut_budget).strip()
                budget_lu = _parse_value(cellule_budget, "float") if cellule_budget else None

                budget_ecarte = (
                    ImportError(
                        row=idx + 1,
                        column="budget",
                        message=f"Budget « {cellule_budget} » illisible, non enregistré",
                        data=mapped,
                    )
                    if cellule_budget and budget_lu is None
                    else None
                )

                if existing and update_existing:
                    existing.name = mapped.get("name") or existing.name
                    if mapped.get("description"):
                        existing.description = mapped["description"]
                    # B-1083 : une colonne absente, vide ou inconnue déliait le
                    # projet de son client et le repassait en « active ».
                    if contact_id:
                        existing.contact_id = contact_id
                    if statut_reconnu:
                        existing.status = statut_reconnu
                    if budget_lu is not None:
                        existing.budget = budget_lu
                    if budget_ecarte:
                        result.errors.append(budget_ecarte)
                    if mapped.get("notes"):
                        existing.notes = mapped["notes"]
                    if mapped.get("tags"):
                        tags_lues, signalement_tags = _lire_les_etiquettes(mapped["tags"], idx + 1, mapped)
                        if tags_lues is not None:
                            existing.tags = tags_lues
                        if signalement_tags:
                            result.errors.append(signalement_tags)

                    existing.updated_at = datetime.now(UTC)
                    self.session.add(existing)
                    result.updated += 1

                elif existing and not update_existing:
                    result.skipped += 1

                else:
                    tags_de_la_ligne, signalement_tags = _lire_les_etiquettes(mapped.get("tags"), idx + 1, mapped)
                    project = Project(
                        id=project_id or generate_uuid(),
                        name=mapped.get("name", "Sans nom"),
                        description=mapped.get("description"),
                        contact_id=contact_id,
                        status=status,
                        budget=budget_lu,
                        notes=mapped.get("notes"),
                        tags=tags_de_la_ligne,
                        scope="global",
                    )
                    self.session.add(project)
                    result.created += 1
                    if signalement_tags:
                        result.errors.append(signalement_tags)
                    if budget_ecarte:
                        result.errors.append(budget_ecarte)

            except Exception as e:
                logger.error(f"Error importing project row {idx + 1}: {e}")
                result.errors.append(ImportError(row=idx + 1, column=None, message=str(e)))
                result.skipped += 1

        await self.session.commit()

        result.success = len(result.errors) == 0
        logger.info(f"Project import: {result.message}")

        return result

    async def import_deliverables(
        self,
        content: bytes,
        filename: str | None = None,
        custom_mapping: dict[str, str] | None = None,
        update_existing: bool = True,
    ) -> ImportResult:
        """
        Import deliverables from file.

        Args:
            content: File content as bytes
            filename: Original filename for format detection
            custom_mapping: Custom column mapping to override defaults
            update_existing: Whether to update existing deliverables by ID

        Returns:
            ImportResult with counts and errors
        """
        format_type = _detect_format(content, filename)
        mapping = {**DELIVERABLE_COLUMN_MAPPING, **(custom_mapping or {})}

        try:
            if format_type == "csv":
                raw_data = _parse_csv(content)
            elif format_type == "xlsx":
                raw_data = _parse_xlsx(content, sheet_name="Livrables")
            else:
                raw_data = _parse_json(content)
        except Exception as e:
            return ImportResult(
                success=False,
                errors=[ImportError(row=0, column=None, message=str(e))],
            )

        result = ImportResult(success=True, total_rows=len(raw_data))

        for idx, row in enumerate(raw_data):
            try:
                mapped = _map_columns(row, mapping)

                errors = _validate_deliverable(mapped)
                if errors:
                    result.errors.append(ImportError(
                        row=idx + 1,
                        column=None,
                        message="; ".join(errors),
                        data=mapped,
                    ))
                    result.skipped += 1
                    continue

                deliverable_id = mapped.get("id")
                existing = None

                if deliverable_id:
                    stmt = select(Deliverable).where(Deliverable.id == deliverable_id)
                    db_result = await self.session.execute(stmt)
                    existing = db_result.scalar_one_or_none()

                # Validate project_id exists
                project_id = mapped.get("project_id")
                if project_id:
                    stmt = select(Project).where(Project.id == project_id)
                    db_result = await self.session.execute(stmt)
                    if not db_result.scalar_one_or_none():
                        result.errors.append(ImportError(
                            row=idx + 1,
                            column="project_id",
                            message=f"Projet {project_id} non trouve",
                            data=mapped,
                        ))
                        result.skipped += 1
                        continue

                raw_status = str(mapped.get("status") or "").lower().strip()
                # B-1320 : même table et même clé que la synchro tableur.
                statut_reconnu = DELIVERABLE_STATUS_MAP.get(cle_de_statut(raw_status))
                status = statut_reconnu or "a_faire"
                # B-1307 : inconnu, le statut ne remplace rien (B-1083, B-1106)
                # et prend le défaut à la création ; il se dit au rapport.
                if raw_status and not statut_reconnu and not (existing and not update_existing):
                    result.errors.append(ImportError(
                        row=idx + 1,
                        column="status",
                        message=f"Statut « {mapped.get('status')} » inconnu, non enregistré",
                        data=mapped,
                    ))

                if existing and update_existing:
                    existing.title = mapped.get("title") or existing.title
                    if mapped.get("description"):
                        existing.description = mapped["description"]
                    existing.project_id = project_id
                    # B-1106 : jumeau de B-1083, un statut absent, vide ou
                    # inconnu repassait le livrable en « a_faire ».
                    if statut_reconnu:
                        existing.status = statut_reconnu
                    if mapped.get("due_date"):
                        # B-1302 : illisible, elle ne remplace pas l'échéance et se dit.
                        echeance = _parse_value(mapped["due_date"], "datetime")
                        if echeance is None:
                            result.errors.append(ImportError(
                                row=idx + 1,
                                column="due_date",
                                message=f"Date « {mapped['due_date']} » illisible, non enregistrée",
                                data=mapped,
                            ))
                        else:
                            existing.due_date = echeance

                    existing.updated_at = datetime.now(UTC)
                    self.session.add(existing)
                    result.updated += 1

                elif existing and not update_existing:
                    result.skipped += 1

                else:
                    echeance_lue = _parse_value(mapped.get("due_date"), "datetime")
                    deliverable = Deliverable(
                        id=deliverable_id or generate_uuid(),
                        title=mapped.get("title", "Sans titre"),
                        description=mapped.get("description"),
                        project_id=project_id,
                        status=status,
                        due_date=echeance_lue,
                    )
                    self.session.add(deliverable)
                    result.created += 1
                    # B-1321 : à la création aussi, une échéance illisible se dit.
                    if echeance_lue is None and str(mapped.get("due_date") or "").strip():
                        result.errors.append(ImportError(
                            row=idx + 1,
                            column="due_date",
                            message=f"Date « {str(mapped['due_date']).strip()} » illisible, non enregistrée",
                            data=mapped,
                        ))

            except Exception as e:
                logger.error(f"Error importing deliverable row {idx + 1}: {e}")
                result.errors.append(ImportError(row=idx + 1, column=None, message=str(e)))
                result.skipped += 1

        await self.session.commit()

        result.success = len(result.errors) == 0
        logger.info(f"Deliverable import: {result.message}")

        return result
